from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path('báo cáo/tài liệu/Báo cáo tuần 5 - Security Flow và Postman Auth Tests.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(10.5)
for name, size, color in [('Title', 20, '17365D'), ('Heading 1', 14, '1F4E79'), ('Heading 2', 11, '2F75B5')]:
    doc.styles[name].font.name = 'Arial'; doc.styles[name].font.size = Pt(size); doc.styles[name].font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill='D9EAF7'):
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(shd)

def code(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25)
    r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(9); return p

p = doc.add_heading('BÁO CÁO TUẦN 5', 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('JWT Security — Security Flow, Postman Auth Tests và Secure Endpoints'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('')
doc.add_heading('1. Nội dung thực hiện', 1)
for x in ['Thiết lập security flow cho API bằng JWT.', 'Kiểm thử authentication flow bằng Postman.', 'Demo secure endpoints với access token.', 'Kiểm tra phân quyền theo role admin/user.']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('2. Security flow', 1)
code('Client → POST /auth/login → AuthController\nAuthController → AuthService → kiểm tra User\nAuthService → tạo JWT access token → Client\nClient → Authorization: Bearer <token> → Secure Endpoint\nAuthentication Middleware → kiểm tra token\nAuthorization Middleware → kiểm tra role\nEndpoint → trả response')
doc.add_paragraph('JWT được dùng để xác thực request. Access token chứa các claim cơ bản như username và role. Server kiểm tra chữ ký, thời hạn và quyền trước khi cho phép truy cập endpoint.')

doc.add_heading('3. Demo secure endpoints', 1)
doc.add_heading('3.1. Đăng nhập lấy token', 2)
code('POST /auth/login\nContent-Type: application/json\n\n{\n  "username": "admin",\n  "password": "123456"\n}')
doc.add_paragraph('Kết quả mong đợi: HTTP 200 và response chứa accessToken.')
code('{\n  "accessToken": "<jwt-token>",\n  "username": "admin",\n  "role": "admin"\n}')
doc.add_heading('3.2. Gọi endpoint yêu cầu đăng nhập', 2)
code('GET /api/profile\nAuthorization: Bearer <accessToken>')
doc.add_paragraph('Token hợp lệ thì API trả thông tin profile của người dùng.')
doc.add_heading('3.3. Gọi endpoint chỉ dành cho admin', 2)
code('GET /api/admin/users\nAuthorization: Bearer <admin-accessToken>')
doc.add_paragraph('Admin nhận HTTP 200. User thường nhận HTTP 403 Forbidden.')

doc.add_heading('4. Postman authentication tests', 1)
t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, text in zip(t.rows[0].cells, ['Test case', 'Request', 'Expected result']): c.text = text; shade(c)
tests = [
    ('Login đúng thông tin', 'POST /auth/login', '200 OK + accessToken'),
    ('Sai password', 'POST /auth/login', '401 Unauthorized'),
    ('Không gửi token', 'GET /api/profile', '401 Unauthorized'),
    ('Token không hợp lệ/hết hạn', 'GET /api/profile', '401 Unauthorized'),
    ('Student gọi API admin', 'GET /api/admin/users', '403 Forbidden'),
    ('Admin gọi API admin', 'GET /api/admin/users', '200 OK'),
]
for row in tests:
    cells = t.add_row().cells
    for c, text in zip(cells, row): c.text = text

doc.add_heading('5. Kết quả đạt được', 1)
for x in ['API đăng nhập trả về JWT access token.', 'Endpoint /api/profile được bảo vệ bằng [Authorize].', 'Endpoint /api/admin/users được bảo vệ theo role admin.', 'Postman kiểm tra được các trường hợp thành công và thất bại.', 'Security flow từ login đến secure endpoint đã được xác nhận.']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('6. Kết luận', 1)
doc.add_paragraph('Tuần 5 đã triển khai và kiểm thử thành công quy trình xác thực JWT, phân quyền theo role và truy cập các secure endpoints. Các kết quả kiểm thử trong Postman phù hợp với mã HTTP kỳ vọng: 200, 401 và 403.')
doc.save(out)
print(out)
