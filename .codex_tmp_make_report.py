from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

d=Document()
d.styles['Normal'].font.name='Arial'; d.styles['Normal'].font.size=Pt(11)
t=d.add_heading('TỔNG HỢP QUÁ TRÌNH HỌC TẬP VÀ THỰC HIỆN PROJECT',0); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=d.add_paragraph('Phạm vi: Tuần 1 – Tuần 3 (W1–W3)'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading('1. Mục tiêu chung',1)
d.add_paragraph('Xây dựng từng bước một ứng dụng quản lý đăng ký khóa học bằng .NET, từ project ASP.NET Core tối giản đến domain/service, unit test, Dependency Injection và repository in-memory.')
d.add_heading('2. Tổng quan kết quả theo tuần',1)
items=[
('W1 – Khởi tạo project ASP.NET Core','Tạo project w1Api trên ASP.NET Core 8. Cấu hình Program.cs với WebApplication builder và endpoint GET “/” trả về “Hello World!”. Bổ sung README hướng dẫn restore, build và run project.'),
('W2 – Domain, service và unit test','Tạo các model Student, Course, Enrollment và Score; định nghĩa repository interface. Xây dựng EnrollmentService.RegisterStudentToCourse: tìm học viên/khóa học, kiểm tra lỗi, tạo Enrollment, thêm vào Student và gọi Save. Viết test xUnit/Moq cho trường hợp hợp lệ và không tìm thấy học viên.'),
('W3 – Repository in-memory và Dependency Injection','Tổ chức Domain/Infrastructure; hoàn thiện service với InvalidOperationException; tạo InMemoryStudentRepository và InMemoryCourseRepository; đăng ký service và repository bằng AddApplicationServices với lifetime Scoped.')
]
for h,x in items: d.add_heading(h,2); d.add_paragraph(x)
d.add_heading('3. Kiến thức và bài học rút ra',1)
for x in ['Cấu trúc và vòng đời khởi động của ASP.NET Core.','Phân tách trách nhiệm giữa model, interface, service và repository.','Dependency Inversion: service phụ thuộc interface, không phụ thuộc cách lưu trữ.','Unit test theo Arrange–Act–Assert; dùng Moq để giả lập repository và Verify lời gọi Save.','Xử lý rõ các nhánh lỗi khi học viên hoặc khóa học không tồn tại.','DI giúp thay implementation repository mà ít ảnh hưởng nghiệp vụ.']: d.add_paragraph(x,style='List Bullet')
d.add_heading('4. Luồng nghiệp vụ chính',1)
d.add_paragraph('RegisterStudentToCourse(studentId, courseId) → tìm Student → kiểm tra → tìm Course → kiểm tra → tạo Enrollment với ngày đăng ký → thêm vào Student → Save → trả về Enrollment.')
d.add_heading('5. Những phần đã hoàn thành',1)
for x in ['Project W1 chạy được với endpoint cơ bản.','Domain model và repository interface đã được tạo.','Nghiệp vụ đăng ký học viên vào khóa học đã được hiện thực.','Có unit test cho luồng thành công và lỗi không tìm thấy học viên.','Có repository in-memory và cấu hình DI ở W3.','Có báo cáo riêng cho từng tuần trong thư mục báo cáo/tài liệu.']: d.add_paragraph(x,style='List Bullet')
d.add_heading('6. Hạn chế và hướng phát triển',1)
d.add_paragraph('Dữ liệu hiện ở mức in-memory; chưa có database thật, API CRUD đầy đủ, validation nâng cao, xử lý trùng đăng ký hoặc kiểm thử toàn diện. Hướng tiếp theo: thêm controller/API, EF Core, logging, global exception handling và mở rộng test coverage.')
d.add_heading('7. Kết luận',1)
d.add_paragraph('Qua W1–W3, project đã đi từ ứng dụng ASP.NET Core tối giản đến cấu trúc có domain service, repository abstraction, unit test và Dependency Injection — nền tảng phù hợp để phát triển ứng dụng quản lý khóa học hoàn chỉnh.')
d.save('báo cáo/Tổng hợp bài học và công việc W1-W3.docx')
print('created')
