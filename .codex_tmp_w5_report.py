from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path('báo cáo/tài liệu/Báo cáo tuần 5 - JWT Security và Demo Postman.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal'].font.size = Pt(10.5)
for name, size, color in [('Title', 22, '17365D'), ('Heading 1', 15, '1F4E79'), ('Heading 2', 12, '2F75B5')]:
    styles[name].font.name = 'Arial'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def code(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.25)
    r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(45,45,45)
    return p

title = doc.add_heading('TUẦN 5 — JWT SECURITY & POSTMAN DEMO', 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('Hướng dẫn học sâu, tự xây dựng và trình bày project w5'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Mục tiêu: hiểu JWT, OAuth2/OIDC cơ bản, OWASP, CORS, Rate Limiting, phân quyền role và tự kiểm thử secure endpoints bằng Swagger/Postman.')

doc.add_heading('1. Kết quả cần đạt', 1)
for x in ['Tự tạo project ASP.NET Core Web API .NET 8.', 'Tạo API đăng nhập và phát hành JWT access token.', 'Bảo vệ endpoint bằng [Authorize].', 'Phân quyền admin/user bằng role claim.', 'Biết kiểm thử các mã 200, 401, 403 và 429.', 'Giải thích được cấu trúc clean code của w5.']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('2. Tạo project và cài package', 1)
doc.add_paragraph('Từ thư mục workspace, tạo project w5 hoặc kế thừa cấu trúc ASP.NET Core của w4:')
code('dotnet new webapi -n w5 --framework net8.0\ncd w5\ndotnet add package Swashbuckle.AspNetCore --version 6.6.2\ndotnet add package Microsoft.AspNetCore.Authentication.JwtBearer --version 8.0.0')
doc.add_paragraph('Sau đó tạo các thư mục: Controllers, Contracts, Domain/Models, Infrastructure và Services.')

doc.add_heading('3. Cấu trúc clean code', 1)
t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
for c, text in zip(t.rows[0].cells, ['Thành phần', 'Trách nhiệm']): c.text = text; shade(c, 'D9EAF7')
for a,b in [('Controllers', 'Nhận HTTP request, gọi service, trả HTTP response.'), ('Contracts', 'DTO input/output của API, validation request.'), ('Domain/Models', 'Đối tượng nghiệp vụ User.'), ('Infrastructure', 'Truy xuất dữ liệu qua repository.'), ('Services', 'Xử lý đăng nhập, tạo claims và JWT.'), ('Program.cs', 'Dependency Injection, middleware và authentication configuration.')]:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b

doc.add_heading('4. Cấu hình JWT', 1)
doc.add_paragraph('Đặt secret và thông tin token trong appsettings.json. Không hard-code secret trong controller/service khi triển khai thật.')
code('{\n  "Jwt": {\n    "Key": "w5-demo-secret-key-change-this-in-production-123456",\n    "Issuer": "w5-api",\n    "Audience": "w5-client"\n  }\n}')
doc.add_paragraph('Program.cs đọc Jwt:Key, đăng ký AddAuthentication().AddJwtBearer(), bật ValidateLifetime, ValidateIssuerSigningKey, Issuer và Audience. Middleware phải theo thứ tự: CORS → RateLimiter → Authentication → Authorization.')

doc.add_heading('5. Luồng đăng nhập JWT', 1)
code('Client -- POST /auth/login --> AuthController\nAuthController --> IAuthService\nAuthService --> IUserRepository tìm User\nNếu đúng --> tạo claims + ký JWT --> trả accessToken\nClient -- Authorization: Bearer <token> --> secure endpoint\nJWT middleware kiểm tra chữ ký/thời hạn --> Authorize kiểm tra role')
doc.add_paragraph('Claims chính: sub là username, Name là username và Role là quyền. Role claim là cơ sở để [Authorize(Roles = "admin")] hoạt động.')

doc.add_heading('6. Demo từng bước bằng Swagger/Postman', 1)
doc.add_heading('Bước 1 — Chạy API', 2)
code('dotnet build w5/w5.csproj\ndotnet run --project w5')
doc.add_paragraph('Mở đường dẫn Swagger được hiển thị trên terminal, thường là https://localhost:xxxx/swagger hoặc http://localhost:xxxx/swagger.')
doc.add_heading('Bước 2 — Đăng nhập admin', 2)
code('POST /auth/login\nContent-Type: application/json\n\n{\n  "username": "admin",\n  "password": "123456"\n}')
doc.add_paragraph('Kết quả mong đợi: 200 OK và response chứa accessToken, username=admin, role=admin. Copy accessToken.')
doc.add_heading('Bước 3 — Gọi profile có token', 2)
code('GET /api/profile\nAuthorization: Bearer <accessToken>')
doc.add_paragraph('Kết quả: 200 OK, trả username và role của người đăng nhập.')
doc.add_heading('Bước 4 — Kiểm tra phân quyền', 2)
code('Đăng nhập student / 123456\nDùng token student gọi GET /api/admin/users\nKết quả mong đợi: 403 Forbidden')
doc.add_paragraph('Đăng nhập admin rồi gọi lại endpoint trên: kết quả mong đợi là 200 OK.')

doc.add_heading('7. Bảng kiểm thử bắt buộc', 1)
t = doc.add_table(rows=1, cols=3); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for c, text in zip(t.rows[0].cells, ['Tình huống', 'Endpoint', 'Kết quả']): c.text=text; shade(c,'D9EAF7')
tests=[('Đăng nhập đúng', 'POST /auth/login', '200'), ('Sai password', 'POST /auth/login', '401'), ('Thiếu token', 'GET /api/profile', '401'), ('Token sai/hết hạn', 'GET /api/profile', '401'), ('User gọi API admin', 'GET /api/admin/users', '403'), ('Admin gọi API admin', 'GET /api/admin/users', '200'), ('Vượt 5 lần login/phút', 'POST /auth/login', '429')]
for row in tests:
    cells=t.add_row().cells
    for c, text in zip(cells,row): c.text=text

doc.add_heading('8. Giải thích các khái niệm phải nắm', 1)
for h, text in [('JWT', 'Token gồm Header.Payload.Signature; payload có thể đọc được nên không đưa password vào token.'), ('OAuth2/OIDC', 'OAuth2 cấp quyền; OIDC bổ sung xác thực danh tính, ví dụ Đăng nhập bằng Google.'), ('CORS', 'Xác định frontend origin nào được phép gọi API.'), ('Rate Limiting', 'Giới hạn số request, đặc biệt hữu ích để chống brute-force login.'), ('401 và 403', '401 là chưa xác thực/token không hợp lệ; 403 là đã xác thực nhưng không đủ quyền.'), ('Clean Code', 'Controller mỏng; nghiệp vụ ở Service; dữ liệu ở Repository; dependency được inject qua interface.')]:
    doc.add_heading(h, 2); doc.add_paragraph(text)

doc.add_heading('9. Lỗi thường gặp và cách xử lý', 1)
for x in ['401 dù đã gửi token: kiểm tra có đúng tiền tố Bearer và token còn hạn không.', '403 với admin: kiểm tra token có ClaimTypes.Role=admin và endpoint có đúng [Authorize(Roles = "admin")] không.', '404 /auth/login: kiểm tra method có [HttpPost("login")] và controller đã được MapControllers chưa.', 'Issuer/Audience invalid: kiểm tra appsettings và giá trị issuer/audience lúc tạo token phải trùng lúc validate.', 'Không gọi được từ frontend: kiểm tra origin đã có trong policy CORS chưa.', 'Password plain text chỉ phục vụ demo; bài production phải hash password bằng BCrypt/ASP.NET Identity.']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('10. Checklist tự tạo lại w5', 1)
for x in ['Tạo csproj và cài JWT package.', 'Tạo User, repository interface và in-memory repository.', 'Tạo LoginRequest/LoginResponse có validation.', 'Viết AuthService kiểm tra user và tạo JWT.', 'Đăng ký DI trong Program.cs.', 'Cấu hình JWT Authentication và Authorization.', 'Viết AuthController và SecureController.', 'Thêm CORS và Rate Limiting.', 'Build project.', 'Test đủ 200/401/403/429 bằng Postman hoặc Swagger.']:
    doc.add_paragraph('☐ ' + x)

doc.add_heading('11. Kết luận', 1)
doc.add_paragraph('W5 hoàn thiện nền tảng bảo mật cho Web API: xác thực bằng JWT, phân quyền role, giới hạn request và kiểm thử security flow. Khi đã tự viết lại được toàn bộ checklist, bước tiếp theo là thay in-memory bằng database, hash password, thêm refresh token, global exception handling, logging và unit/integration test.')

doc.save(out)
print(out)
